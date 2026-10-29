"""Fixed, bounded Office reader; no model code or network. Mirrors DSH Office library choices."""
import json
import sys
import zipfile
from pathlib import Path

source, start, limit = Path(sys.argv[1]), int(sys.argv[2]), int(sys.argv[3])
if start < 1 or not 1 <= limit <= 20:
    raise ValueError("invalid window")
with zipfile.ZipFile(source) as archive:
    entries = archive.infolist()
    if len(entries) > 10000 or sum(item.file_size for item in entries) > 128 * 1024 * 1024:
        raise ValueError("archive limits")
    if any(item.flag_bits & 1 for item in entries):
        raise ValueError("encrypted archive")

items, warnings, truncated = [], [], False

def add(index, parts):
    global truncated
    text = "\n".join(parts)
    if len(text) > 40000:
        truncated = True
    items.append({"index": index, "text": text[:40000], "truncated": len(text) > 40000})

if source.suffix == ".docx":
    from docx import Document
    document = Document(source)
    units = list(document.iter_inner_content())
    for index in range(start - 1, min(len(units), start - 1 + limit)):
        block = units[index]
        if hasattr(block, "rows"):
            add(index + 1, [" | ".join(cell.text for cell in row.cells) for row in block.rows])
        else:
            add(index + 1, [block.text])
    unit, total = "block", len(units)
    warnings = ["读取正文段落与表格；未读取页眉页脚、批注及嵌入图片。Word分页需另行渲染。"]
elif source.suffix == ".pptx":
    from pptx import Presentation
    presentation = Presentation(source)
    for index in range(start - 1, min(len(presentation.slides), start - 1 + limit)):
        slide, parts = presentation.slides[index], []
        def read_shapes(shapes):
            for shape in shapes:
                if hasattr(shape, "shapes"):
                    read_shapes(shape.shapes)
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                if shape.has_table:
                    parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                if shape.has_chart:
                    parts.append("[图表：结构读取未包含完整视觉信息，请按需渲染该页]")
                    for series in shape.chart.series:
                        parts.append(str(series.name) + ": " + ", ".join(str(v) for v in series.values))
        read_shapes(slide.shapes)
        if slide.has_notes_slide:
            parts.append("[备注]\n" + slide.notes_slide.notes_text_frame.text)
        add(index + 1, parts)
    unit, total = "slide", len(presentation.slides)
    warnings = ["读取文字、表格、可访问图表数值及备注；未理解图片、动画或SmartArt版面，需按需渲染。"]
else:
    raise ValueError("unsupported format")
print(json.dumps({"unit": unit, "total": total, "items": items, "warnings": warnings, "truncated": truncated}, ensure_ascii=False))
